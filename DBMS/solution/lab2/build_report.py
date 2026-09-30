"""Сборка отчёта по ЛР 2 (СУБД): python build_report.py -> report.docx."""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt

STUDENT = "ФИО"
GROUP = "ФИТ-242"
COURSE = "3"
DIRECTION = "02.03.02 Фундаментальная информатика и информационные технологии"
TEACHER = "Моисеева Н.А."

HERE = Path(__file__).resolve().parent
IMG = HERE / "img"

TABLES = [
    ("teams", "Команды", [
        ("team_ID", "integer", "PK, NOT NULL", "Идентификатор команды"),
        ("name", "varchar(100)", "NOT NULL, UNIQUE", "Название команды"),
        ("base", "varchar(100)", "NOT NULL", "База (город) команды"),
    ]),
    ("cars", "Болиды", [
        ("car_ID", "integer", "PK, NOT NULL", "Идентификатор болида"),
        ("model", "varchar(50)", "NOT NULL", "Модель"),
        ("year", "integer", "NOT NULL, CHECK 1950–2100", "Год выпуска"),
        ("team_ID", "integer", "FK → teams, NOT NULL", "Команда"),
    ]),
    ("drivers", "Пилоты", [
        ("driver_ID", "integer", "PK, NOT NULL", "Идентификатор пилота"),
        ("name", "varchar(100)", "NOT NULL", "Имя пилота"),
        ("country", "varchar(50)", "NOT NULL", "Страна"),
        ("team_ID", "integer", "FK → teams, NOT NULL", "Команда"),
    ]),
    ("parts", "Запчасти", [
        ("car_ID", "integer", "PK, FK → cars, NOT NULL", "Болид"),
        ("part_ID", "integer", "PK, NOT NULL", "Номер запчасти в болиде"),
        ("name", "varchar(100)", "NOT NULL", "Название"),
        ("type", "varchar(50)", "NOT NULL", "Тип запчасти"),
    ]),
    ("sponsors", "Спонсоры", [
        ("sponsor_ID", "integer", "PK, NOT NULL", "Идентификатор спонсора"),
        ("name", "varchar(100)", "NOT NULL, UNIQUE", "Название"),
        ("budget", "numeric(12,2)", "NOT NULL, CHECK ≥ 0", "Бюджет"),
    ]),
    ("tracks", "Трассы", [
        ("track_ID", "integer", "PK, NOT NULL", "Идентификатор трассы"),
        ("name", "varchar(100)", "NOT NULL, UNIQUE", "Название"),
        ("country", "varchar(50)", "NOT NULL", "Страна"),
        ("length_km", "numeric(5,3)", "NOT NULL, CHECK > 0", "Длина круга, км"),
    ]),
    ("races", "Гонки", [
        ("race_ID", "integer", "PK, NOT NULL", "Идентификатор гонки"),
        ("name", "varchar(100)", "NOT NULL", "Название Гран-при"),
        ("date", "date", "NOT NULL", "Дата гонки"),
        ("track_ID", "integer", "FK → tracks, NOT NULL", "Трасса"),
        ("sponsor_ID", "integer", "FK → sponsors, NOT NULL", "Спонсор"),
    ]),
]

FOREIGN_KEYS = [
    ("cars_FK", "cars.team_ID → teams.team_ID", "1:N", "CASCADE", "CASCADE"),
    ("drivers_FK", "drivers.team_ID → teams.team_ID", "1:N", "CASCADE", "CASCADE"),
    ("parts_FK", "parts.car_ID → cars.car_ID", "1:N, идентифицирующая", "CASCADE", "CASCADE"),
    ("races_FK1", "races.track_ID → tracks.track_ID", "1:N", "RESTRICT", "CASCADE"),
    ("races_FK2", "races.sponsor_ID → sponsors.sponsor_ID", "1:N", "RESTRICT", "CASCADE"),
]

INDEXES = [
    ("teams_PK, cars_PK, drivers_PK, sponsors_PK, tracks_PK, races_PK", "первичный ключ таблицы",
     "автоматически (PRIMARY KEY)"),
    ("parts_PK", "(car_ID, part_ID)", "автоматически (PRIMARY KEY); покрывает и внешний ключ parts.car_ID"),
    ("teams_name_UQ, sponsors_name_UQ, tracks_name_UQ", "name", "автоматически (UNIQUE)"),
    ("cars_team_IDX", "cars.team_ID", "CREATE INDEX"),
    ("drivers_team_IDX", "drivers.team_ID", "CREATE INDEX"),
    ("races_track_IDX", "races.track_ID", "CREATE INDEX"),
    ("races_sponsor_IDX", "races.sponsor_ID", "CREATE INDEX"),
]

CHECKS = [
    ("Вставка болида с несуществующей командой", "foreign_key_violation (cars_FK)"),
    ("Повтор пары (car_ID, part_ID) в parts", "unique_violation (parts_PK)"),
    ("Повтор названия команды", "unique_violation (teams_name_UQ)"),
    ("Болид 1900 года", "check_violation (cars_year_CHK)"),
    ("Трасса длиной 0 км", "check_violation (tracks_length_CHK)"),
    ("Спонсор с отрицательным бюджетом", "check_violation (sponsors_budget_CHK)"),
    ("Пилот без имени", "not_null_violation"),
    ("Удаление трассы, на которой есть гонки", "foreign_key_violation (races_FK1, RESTRICT)"),
    ("Удаление спонсора, у которого есть гонки", "foreign_key_violation (races_FK2, RESTRICT)"),
    ("UPDATE teams: team_ID 1 → 10", "в cars и drivers team_ID стал 10 (ON UPDATE CASCADE)"),
    ("DELETE команды 10", "удалены её болид, пилот и запчасти болида (ON DELETE CASCADE)"),
]

FIGURES = [
    ("01_database.png", "База данных formula1 в браузере pgAdmin и её SQL-определение"),
    ("02_tables.png", "Таблицы базы данных formula1"),
    ("03_constraints.png", "Таблица cars: первичный и внешний ключ, ограничение CHECK, индекс"),
    ("07_parts_pk.png", "Таблица parts: составной первичный ключ и идентифицирующая связь с cars"),
    ("04_indexes.png", "Таблица races: два внешних ключа с RESTRICT и индексы на них"),
    ("05_check_messages.png", "Результат выполнения 03_check.sql (вкладка Messages)"),
    ("06_erd.png", "ER-диаграмма базы данных formula1 (ERD Tool → ERD for Database)"),
]


def para(doc, text="", bold=False, align=None, size=14, indent=True, space_after=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.first_line_indent = Cm(1.25) if indent and align is None else Cm(0)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = 1.5 if size >= 14 else 1.0
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
    return p


def heading(doc, text):
    p = para(doc, text, bold=True, indent=False, space_after=6)
    p.paragraph_format.keep_with_next = True
    return p


def code(doc, text):
    for line in text.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0
        r = p.add_run(line if line else " ")
        r.font.name = "Courier New"
        r.font.size = Pt(9)


def table(doc, header, rows, widths=None, size=11):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, h in zip(t.rows[0].cells, header):
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for cell, v in zip(cells, row):
            cell.text = ""
            cell.paragraphs[0].add_run(v).font.size = Pt(size)
    if widths:
        for row in t.rows:
            for cell, w in zip(row.cells, widths):
                cell.width = Cm(w)
    doc.add_paragraph()
    return t


def figure(doc, n, name, caption):
    path = IMG / name
    if not path.exists():
        para(doc, f"[Рисунок {n} — вставить скриншот img/{name}]", align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(16.5))
    para(doc, f"Рисунок {n} — {caption}", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=6)


def title_page(doc):
    c = WD_ALIGN_PARAGRAPH.CENTER
    for line in (
        "Министерство науки и высшего образования РФ",
        "федеральное государственное автономное образовательное учреждение высшего образования",
        "«Омский государственный технический университет»",
    ):
        para(doc, line, align=c)
    para(doc)
    para(doc, "Факультет информационных технологий и компьютерных систем", align=c)
    para(doc, "Кафедра «Прикладная математика и фундаментальная информатика»", align=c)
    for _ in range(4):
        para(doc)
    para(doc, "Лабораторная работа № 2", bold=True, align=c, size=16)
    para(doc, "по дисциплине «Системы управления базами данных»", align=c)
    para(doc, "Тема: Создание БД, таблиц, определение индексов и связей в PostgreSQL", align=c)
    para(doc, "Предметная область: чемпионат «Формула-1»", align=c)
    for _ in range(3):
        para(doc)
    rows = [
        ("Студента", STUDENT),
        ("Курс", COURSE),
        ("Группа", GROUP),
        ("Направление", DIRECTION),
        ("Преподаватель", TEACHER),
        ("Выполнил", "дата, подпись студента"),
        ("Проверил", "дата, подпись преподавателя"),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].width = Cm(4)
        cells[1].width = Cm(9)
        cells[0].paragraphs[0].add_run(k).font.size = Pt(14)
        cells[1].paragraphs[0].add_run(v).font.size = Pt(14)
    for _ in range(4):
        para(doc)
    p = para(doc, "Омск 2026", align=c)
    p.runs[0].add_break(WD_BREAK.PAGE)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(14)

    title_page(doc)

    heading(doc, "Цель работы")
    para(doc, "Освоить создание базы данных, таблиц, ограничений целостности, связей и индексов "
              "в СУБД PostgreSQL 15 средствами SQL и pgAdmin 4 на основе модели данных, "
              "спроектированной в лабораторной работе № 1.")

    heading(doc, "Задание")
    para(doc, "Для заданной предметной области средствами PostgreSQL необходимо (методика рейтингового "
              "контроля, ЛР 2; учебное пособие «Работа с СУБД PostgreSQL», с. 68, 76):")
    for item in (
        "создать новую базу данных;",
        "создать таблицы, определить поля таблиц и тип полей;",
        "определить связи между таблицами и ограничения целостности;",
        "определить индексы;",
        "составить отчёт.",
    ):
        para(doc, "• " + item)
    para(doc, "Предметная область — чемпионат «Формула-1». Модель данных (ЛР 1) содержит семь сущностей: "
              "Команды, Болиды, Пилоты, Запчасти, Спонсоры, Трассы и Гонки. Команда имеет много болидов "
              "и пилотов, болид — много запчастей (запчасть — зависимая сущность, идентифицирующая связь), "
              "на трассе и при участии спонсора проводится много гонок.")

    heading(doc, "1. Создание базы данных")
    para(doc, "База данных formula1 создаётся инструкцией CREATE DATABASE по образцу пособия (с. 69): "
              "владелец postgres, кодировка UTF8, табличное пространство pg_default, без ограничения "
              "числа подключений. Скрипт выполняется в Query Tool при подключении к базе postgres "
              "(файл 01_create_database.sql).")
    code(doc, (HERE / "01_create_database.sql").read_text(encoding="utf-8"))
    para(doc)
    para(doc, "В файле 01_create_database.sql одна команда: pgAdmin выполняет несколько команд "
              "одной транзакцией, а CREATE DATABASE внутри транзакции запрещён (ошибка SQLSTATE 25001). "
              "Поэтому комментарий к базе задаётся в начале файла 02_create_tables.sql:")
    sql2 = (HERE / "02_create_tables.sql").read_text(encoding="utf-8")
    c0 = sql2.index("COMMENT ON DATABASE")
    code(doc, sql2[c0:sql2.index(";", c0) + 1])
    para(doc)
    para(doc, "Локаль Russian_Russia.1251 указана как в образце и существует в Windows. Скриншоты в отчёте "
              "получены на тестовом сервере PostgreSQL 15 под Linux, где этой локали нет, поэтому на "
              "рисунке 1 показана локаль сервера en_US.utf8.")
    figure(doc, 1, *FIGURES[0])

    heading(doc, "2. Создание таблиц")
    para(doc, "По физической модели из ЛР 1 созданы семь таблиц. Как рекомендует пособие (с. 76), "
              "в инструкциях CREATE TABLE заданы только ограничения NOT NULL и PRIMARY KEY, остальные "
              "ограничения добавляются отдельными инструкциями ALTER TABLE (файл 02_create_tables.sql). "
              "Структура таблиц приведена в таблицах 1–7.")
    for i, (name, ru, cols) in enumerate(TABLES, 1):
        para(doc, f"Таблица {i} — {name} ({ru})", indent=False, size=12)
        table(doc, ("Поле", "Тип", "Ограничения", "Описание"), cols, widths=(3, 3.2, 5, 5.3))
    sql = (HERE / "02_create_tables.sql").read_text(encoding="utf-8")
    fk = sql.index("ALTER TABLE cars\nADD CONSTRAINT cars_FK")
    first = sql.index("CREATE TABLE teams")
    code(doc, sql[first:fk])
    para(doc)
    figure(doc, 2, *FIGURES[1])

    heading(doc, "3. Связи и ограничения целостности")
    para(doc, "Связи «один-ко-многим» из модели ЛР 1 реализованы внешними ключами (таблица 8). "
              "Для данных, принадлежащих команде и болиду, выбрано каскадное удаление и обновление, как "
              "в образце пособия (с. 79): при удалении команды удаляются её болиды и пилоты, при удалении "
              "болида — его запчасти. Трассу и спонсора, у которых есть гонки, удалить нельзя (RESTRICT), "
              "чтобы не потерять результаты гонок; изменение их ключа распространяется каскадно.")
    para(doc, "Таблица 8 — Внешние ключи", indent=False, size=12)
    table(doc, ("Ограничение", "Связь", "Тип", "ON DELETE", "ON UPDATE"), FOREIGN_KEYS,
          widths=(2.8, 6, 3.2, 2.3, 2.2))
    para(doc, "Дополнительно заданы ограничения уникальности потенциальных ключей — названий команд, "
              "спонсоров и трасс (teams_name_UQ, sponsors_name_UQ, tracks_name_UQ) — и ограничения "
              "проверки: год болида от 1950 до 2100 (cars_year_CHK), неотрицательный бюджет спонсора "
              "(sponsors_budget_CHK), положительная длина трассы (tracks_length_CHK).")
    start = sql.index("ALTER TABLE cars\nADD CONSTRAINT cars_FK")
    end = sql.index("CREATE INDEX")
    code(doc, sql[start:end])
    para(doc)
    figure(doc, 3, *FIGURES[2])
    figure(doc, 4, *FIGURES[3])

    heading(doc, "4. Индексы")
    para(doc, "PostgreSQL автоматически создаёт уникальный индекс для каждого первичного ключа и "
              "ограничения UNIQUE (пособие, с. 72), но не создаёт индекс для внешнего ключа (с. 73). "
              "Поэтому на столбцы внешних ключей созданы индексы B-дерева; они ускоряют соединение "
              "таблиц и проверку ссылочной целостности при удалении и изменении родительской строки. "
              "Для parts.car_ID отдельный индекс не нужен: этот столбец — первый в составном "
              "первичном ключе (car_ID, part_ID). Всего в базе 14 индексов (таблица 9).")
    para(doc, "Таблица 9 — Индексы", indent=False, size=12)
    table(doc, ("Индекс", "Столбцы", "Как создан"), INDEXES, widths=(6.5, 4.5, 5.5))
    code(doc, sql[end:sql.index("COMMENT ON TABLE")])
    para(doc)
    figure(doc, 5, *FIGURES[4])

    heading(doc, "5. Проверка")
    para(doc, "Скрипт 03_check.sql выводит из системного каталога список таблиц, полей, ограничений "
              "и индексов, а затем внутри транзакции вставляет тестовые данные и пытается нарушить "
              "каждое ограничение. В конце выполняется ROLLBACK, поэтому база остаётся пустой. "
              "Результаты приведены в таблице 10 и на рисунке 6, полный вывод — в файле check_output.txt.")
    para(doc, "Таблица 10 — Проверка ограничений", indent=False, size=12)
    table(doc, ("Действие", "Результат"), CHECKS, widths=(8, 8.5))
    figure(doc, 6, *FIGURES[5])

    heading(doc, "6. ER-диаграмма")
    para(doc, "ER-диаграмма построена в pgAdmin командой Tools > ERD Tool, пункт ERD for Database "
              "(пособие, с. 80). Она совпадает с моделью ЛР 1: семь таблиц и пять связей "
              "«один-ко-многим».")
    figure(doc, 7, *FIGURES[6])

    heading(doc, "Вывод")
    para(doc, "В ходе работы в PostgreSQL 15 создана база данных formula1 из семи таблиц. Определены типы "
              "полей, первичные ключи (в том числе составной ключ зависимой таблицы parts), пять внешних "
              "ключей с действиями ON DELETE и ON UPDATE, ограничения NOT NULL, UNIQUE и CHECK, а также "
              "индексы на столбцы внешних ключей. Проверка на тестовых данных показала, что все "
              "ограничения отклоняют некорректные данные, а каскадные действия работают согласно "
              "модели предметной области.")

    doc.save(HERE / "report.docx")
    print("report.docx собран")


if __name__ == "__main__":
    main()
